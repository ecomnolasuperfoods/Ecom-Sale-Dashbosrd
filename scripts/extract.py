"""แปลงไฟล์ Excel ต้นทางของ Nola Superfoods เป็น data/data.json สำหรับแดชบอร์ด

วิธีใช้:
    python3 scripts/extract.py [path/to/workbook.xlsx]

ถ้าไม่ระบุ path จะไล่หาไฟล์ .xlsx ไฟล์แรกในโฟลเดอร์ราก
"""
import openpyxl, json, math, datetime, collections, os, sys, glob

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if len(sys.argv) > 1:
    SRC = sys.argv[1]
else:
    cands = sorted(glob.glob(os.path.join(ROOT, '*.xlsx')))
    if not cands:
        sys.exit('ไม่พบไฟล์ .xlsx ในโฟลเดอร์ ' + ROOT + ' — ระบุ path เป็นอาร์กิวเมนต์ได้')
    SRC = cands[0]
OUT = os.path.join(ROOT, 'data', 'data.json')
print('อ่านจาก:', SRC)

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
MON = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']

def num(v):
    if v is None: return None
    if isinstance(v,(int,float)):
        if isinstance(v,float) and (math.isnan(v) or math.isinf(v)): return None
        return round(v,4)
    s=str(v).strip().replace(',','')
    if s in ('','-','#DIV/0!','#N/A',' '): return None
    try: return round(float(s),4)
    except: return None

out={}
out['meta']={'source':os.path.basename(SRC),'generated':datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}

# ---------- 1. YEARLY ----------
ws=wb['Sale Yealy']
rows=[list(r) for r in ws.iter_rows(max_row=89,max_col=7,values_only=True)]
yearly=collections.defaultdict(list)
cur=None
for r in rows:
    p=r[0]
    if p in (None,''): continue
    if p=='Platform': continue
    if p in ('Attribute','YoY diff'): continue
    y=num(r[1])
    if y is None: continue
    yearly[p].append({'year':int(y),'gmv':num(r[3]),'orders':num(r[4]),'aov':num(r[5]),'clicks':num(r[6])})
out['yearly']=dict(yearly)

# ---------- 2. MONTHLY 2026 + ADS ----------
# ชื่อชีตเปลี่ยนตามเวอร์ชันไฟล์ (v14 = '... + Ads Spend', v16 = '... + Ads 2026')
ADS_SHEET=None
for cand in ('Pivot Sale Monthly + Ads 2026','Pivot Sale Monthly + Ads Spend'):
    if cand in wb.sheetnames: ADS_SHEET=cand; break
if ADS_SHEET is None: sys.exit('ไม่พบชีตยอดขายรายเดือน + Ads')
ws=wb[ADS_SHEET]
g=[list(r) for r in ws.iter_rows(max_row=40,max_col=119,values_only=True)]
# หาคอลัมน์เริ่มของบล็อกแต่ละแพลตฟอร์มจากหัวตาราง (แถว index 3) แทนการ hardcode
hdr3=g[3]
hdr4=g[4]
def block_start(label):
    # ค้นเฉพาะฝั่งขวาของตาราง (คอลัมน์ 14 ขึ้นไป) เพราะฝั่งซ้ายเป็น Column Labels ของ pivot
    for ci in range(14,len(hdr3)):
        v=hdr3[ci]
        if v and str(v).strip()==label and str(hdr4[ci] or '').strip()=='GMV': return ci
    return None
BLK={'Lazada':block_start('Lazada'),'Shopee':block_start('Shopee'),'Tiktok':block_start('Tiktok')}
if None in BLK.values(): sys.exit('หาคอลัมน์บล็อกแพลตฟอร์มไม่เจอ: '+repr(BLK))
print('บล็อกแพลตฟอร์ม:', BLK)
plats=['Lazada','Line Shopping','Shopee','Tiktok','Facebook','Line Chat','Amaze','Website']
monthly=[]
for i,m in enumerate(MON):
    r=g[5+i]
    rec={'month':m,'byPlatform':{p:num(r[1+j]) for j,p in enumerate(plats)},
         'total':num(r[9]),'mom':num(r[10]),'adsSpend':num(r[11]),'adsPct':num(r[12]),'roi':num(r[13])}
    monthly.append(rec)
out['monthly2026']=monthly
# targets rows 19..29 (index)
tg={}
hdr=g[19]
for r in g[19:30]:
    if not r[0]: continue
    key=r[0]
    if key in ('Target',): continue
    tg[key]={MON[k]:num(r[1+k]) for k in range(12)}
    tg[key]['Total']=num(r[13])
out['targets2026']=tg

# platform ads detail (cols 15+)
def block(start, fields):
    res=[]
    for i,m in enumerate(MON):
        r=g[5+i]
        res.append({'month':m, **{f:num(r[start+k]) for k,f in enumerate(fields)}})
    return res
out['adsDetail']={
 'Lazada': block(BLK['Lazada'],['gmv','gmvTarget','adsMaxGmv','adsMaxSpend','roas','adsKwGmv','adsKwSpend','roasKw','adsAffGmv','adsAffSpend','roasAff','roi','diffMoMPct','diffMoM','diffTgPct','diffTg']),
 'Shopee': block(BLK['Shopee'],['gmv','gmvTarget','adsGmv','adsSpend','roas','affGmv','affSpend','roasAff','cpasGmv','cpasSpend','roasCpas','roi','diffMoMPct','diffMoM','diffTgPct','diffTg']),
 'Tiktok': block(BLK['Tiktok'],['gmv','gmvTarget','gmvMaxAllGmv','gmvMaxAllSpend','gmvMaxAllRoas','gmvMaxGmv','gmvMaxSpend','gmvMaxRoas','gmvMaxLiveGmv','gmvMaxLiveSpend','gmvMaxLiveRoas','liveRatio','cAdsSpend']),
}

# ---------- 3. DAILY from GMV Database ----------
ws=wb['GMV Database']
it=ws.iter_rows(values_only=True)
hdr=next(it)
daily=collections.defaultdict(lambda: collections.defaultdict(dict))
plat=year=mon=None
extra=collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(float)))
for r in it:
    if r[0]: plat=str(r[0]).strip()
    if r[1]: year=r[1]
    if r[2]: mon=str(r[2]).strip()
    d=num(r[4])
    if d is None or plat is None or year is None or mon is None: continue
    gmv=num(r[5])
    if gmv is None: continue
    key=f'{plat}|{year}|{mon}'
    daily[key][int(d)]={'gmv':gmv,'orders':num(r[6]),'aov':num(r[7]),'clicks':num(r[8]),'visitors':num(r[9]),'cvr':num(r[10]),
                        'cancelOrders':num(r[11]),'cancelGmv':num(r[12]),'refundOrders':num(r[13]),'refundGmv':num(r[14]),
                        'buyers':num(r[15]),'newBuyers':num(r[16]),'returnBuyers':num(r[17]),
                        'chProduct':num(r[22]),'chLive':num(r[23]),'chVideo':num(r[24]),'chPartner':num(r[25]),'chAds':num(r[26])}
out['daily']={k:{str(d):v for d,v in sorted(vv.items())} for k,vv in daily.items()}

# monthly aggregate 2025 vs 2026 for compare
agg=collections.defaultdict(lambda: collections.defaultdict(float))
for key,days in daily.items():
    p,y,m=key.split('|')
    for d,v in days.items():
        for f in ('gmv','orders','clicks','visitors','cancelGmv','newBuyers','returnBuyers','chProduct','chLive','chVideo','chPartner','chAds'):
            if v.get(f): agg[f'{p}|{y}|{m}'][f]+=v[f]
out['monthlyAgg']={k:{f:round(x,2) for f,x in v.items()} for k,v in agg.items()}

# ---------- 4. PRODUCTS (Mapping all stock) ----------
ws=wb['Mapping all stock']
rows=[list(r) for r in ws.iter_rows(min_row=3,max_row=74,max_col=50,values_only=True)]
hdr=rows[0]
# หาคอลัมน์จากหัวตารางแทน hardcode (ไฟล์เพิ่มคอลัมน์เดือนใหม่ทุกรอบ ตำแหน่งจึงเลื่อน)
MON_FULL={'SEP':'Sep','OCT':'Oct','NOV':'Nov','DEC':'Dec','JAN':'Jan','FEB':'Feb','MAR':'Mar','APR':'Apr','MAY':'May','JUN':'Jun','JUL':'Jul','AUG':'Aug'}
monthCols=[]
import re as _re
for ci,v in enumerate(hdr):
    if not v: continue
    s=' '.join(str(v).split()).strip()
    m=_re.match(r'^ยอดขาย\s*([A-Za-z]{3})\s*(\d{2})$', s)
    if m and m.group(1).upper() in MON_FULL:
        monthCols.append((ci, MON_FULL[m.group(1).upper()]+m.group(2)))
def find(label, frm=25):
    for ci in range(frm,len(hdr)):
        if hdr[ci] and ' '.join(str(hdr[ci]).split()).strip().startswith(label): return ci
    return None
C={'avg3m':find('Sale AVG 3 Month'),'avgDday':find('Sale AVG D-DAY'),
   'bauStatus':find('BAU สถานะ'),'ddayStatus':find('D-DAY'),
   'runScore':find('Run Score'),'stockLBL':None,'stockMHC':find('มหาชัย'),
   'planSend':find('จำนวน'),'status':find('Remark')}
# 'ลาดบัวหลวง' โผล่สองที่ (BAU สถานะ กับ ยอดสต็อก) — เอาตัวที่อยู่หลัง Run Score
rs=C['runScore']
C['stockLBL']=find('ลาดบัวหลวง', rs if rs else 25)
print('คอลัมน์สต็อก:', C, '| เดือน:', [m for _,m in monthCols])
prods=[]
for r in rows[2:]:
    if not r[0] or not r[2]: continue
    qty={name:num(r[ci]) for ci,name in monthCols}
    prods.append({'sku':str(r[0]).strip(),'category':str(r[1] or '').strip(),'name':str(r[2]).strip(),'pack':str(r[3] or '').strip(),
                  'qty':qty,'avg3m':num(r[C['avg3m']]),'avgDday':num(r[C['avgDday']]),
                  'bauStatus':str(r[C['bauStatus']] or '').strip(),'ddayStatus':str(r[C['ddayStatus']] or '').strip(),
                  'runScore':num(r[C['runScore']]),'stockLBL':num(r[C['stockLBL']]),'stockMHC':num(r[C['stockMHC']]),
                  'planSend':num(r[C['planSend']]) if C['planSend'] else None,
                  'status':str(r[C['status']] or '').strip() if C['status'] else ''})
out['products']=prods
out['skuMonths']=[m for _,m in monthCols]

# price
ws=wb['Price']
prices=[]
for r in ws.iter_rows(min_row=5,max_row=104,max_col=18,values_only=True):
    if not r[1]: continue
    prices.append({'sku':str(r[1]).strip(),'type':str(r[2] or ''),'name':str(r[3] or ''),'pack':str(r[4] or ''),
                   'sticker':num(r[5]),'ao':num(r[6]),'mmpd':num(r[8]),'dd':num(r[10]),'mega':num(r[12]),'deepest':num(r[14])})
out['prices']=prices

# ---------- 5. AFFILIATE TIKTOK ----------
ws=wb['Affiliate รวม Database Tiktok']
it=ws.iter_rows(min_row=2,max_col=13,values_only=True)
aff=collections.defaultdict(lambda: collections.defaultdict(float))
creators=collections.defaultdict(lambda: collections.defaultdict(float))
namesByMonth=collections.defaultdict(set)
for r in it:
    m=r[0]; name=r[1]
    if not name: continue
    if isinstance(m,datetime.datetime): mk=m.strftime('%b %y')
    else: mk=str(m)[:12]
    gmv=num(r[2]) or 0; comm=num(r[11]) or 0; orders=num(r[4]) or 0
    aff[mk]['gmv']+=gmv; aff[mk]['commission']+=comm; aff[mk]['orders']+=orders; aff[mk]['creators']+=1
    aff[mk]['videos']+= (num(r[9]) or 0); aff[mk]['lives']+= (num(r[10]) or 0)
    creators[name]['gmv']+=gmv; creators[name]['commission']+=comm; creators[name]['orders']+=orders
    if gmv>0: namesByMonth[mk].add(name)
out['affTiktokMonthly']={k:{f:round(v,2) for f,v in d.items()} for k,d in aff.items()}
for k in out['affTiktokMonthly']: out['affTiktokMonthly'][k]['creators']=len(namesByMonth[k])
top=sorted(creators.items(), key=lambda kv:-kv[1]['gmv'])[:40]
out['affTiktokTop']=[{'name':k,**{f:round(v,2) for f,v in d.items()}} for k,d in top]

# followers
ws=wb['Affiliate List Tiktok']
out['affTiktokFollowers']={str(r[0]).strip():num(r[1]) for r in ws.iter_rows(min_row=2,max_col=2,values_only=True) if r[0]}

# shopee affiliate
ws=wb['Pivot Affiliate Shopee']
sa=[]
for r in ws.iter_rows(min_row=4,max_row=143,max_col=7,values_only=True):
    if not r[0] or str(r[0]).startswith('Grand'): continue
    sa.append({'name':str(r[0]).strip(),'gmv':num(r[1]),'qty':num(r[2]),'commission':num(r[3]),'roi':num(r[4]),'buyers':num(r[5]),'newBuyers':num(r[6])})
out['affShopee']=sa
ws=wb['Pivot Affiliate Shopee']
ch=[]
for r in ws.iter_rows(min_row=4,max_row=8,min_col=9,max_col=12,values_only=True):
    if r[0]: ch.append({'channel':str(r[0]),'gmv':num(r[1]),'qty':num(r[2]),'commission':num(r[3])})
out['affShopeeChannel']=ch

# lazada affiliate
ws=wb['Affiliate Database Lazada']
la=[]
for r in ws.iter_rows(min_row=1,max_row=32,max_col=7,values_only=True):
    la.append([str(c)[:40] if c is not None else '' for c in r])
out['affLazadaRaw']=la

# ---------- 6. REPURCHASE ----------
ws=wb['Shopee Report Repurchase']
rp=[]
for r in ws.iter_rows(min_row=5,max_row=7,max_col=16,values_only=True):
    if r[1]:
        rp.append({'metric':str(r[1]).strip(),**{MON[i]:num(r[2+i]) for i in range(12)},'total':num(r[14])})
out['shopeeRepurchase']=rp
ws=wb['Tiktok Report Repurchase']
tr=[]
for r in ws.iter_rows(min_row=6,max_row=37,max_col=14,values_only=True):
    if r[0] is None: continue
    tr.append({'day':num(r[0]),**{MON[i]:num(r[1+i]) for i in range(12)}})
out['tiktokRepurchaseDaily']=tr

# ---------- 7. D-DAY ----------
ws=wb['Sale D-Day']
dd=[]
for r in ws.iter_rows(min_row=2,max_row=41,max_col=13,values_only=True):
    if not r[2] or num(r[4]) is None: continue
    dd.append({'platform':str(r[0] or ''),'year':num(r[1]),'dday':str(r[2]),'gmv':num(r[4]),'orders':num(r[5]),'aov':num(r[6]),'clicks':num(r[7]),'visitors':num(r[8]),'cvr':num(r[9])})
out['dday']=dd

os.makedirs(os.path.dirname(OUT), exist_ok=True)
json.dump(out, open(OUT,'w'), ensure_ascii=False, separators=(',',':'))
print('เขียน:', OUT, os.path.getsize(OUT), 'bytes')
for k,v in out.items():
    print(k, type(v).__name__, len(v) if hasattr(v,'__len__') else '')
