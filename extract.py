"""แปลงไฟล์ Excel ต้นทางของ Nola Superfoods เป็น data/data.json สำหรับแดชบอร์ด

ใช้ไฟล์ต้นทาง 2 ไฟล์:
  1. ไฟล์ยอดขายหลัก  — ชื่อมีคำว่า "Sale 2026"  (ทุกแท็บยกเว้นสินค้า & สต็อก)
  2. ไฟล์รายงานสินค้า — ชื่อมีคำว่า "Products"   (แท็บสินค้า & สต็อก) · ไม่มีก็ข้ามได้

วิธีใช้:
    python3 scripts/extract.py [main.xlsx] [products.xlsx]

ถ้าไม่ระบุ path จะไล่หาเองจากชื่อไฟล์ .xlsx ในโฟลเดอร์ราก
"""
import openpyxl, json, math, datetime, collections, os, sys, glob, re

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def pick(pattern, taken=()):
    for f in sorted(glob.glob(os.path.join(ROOT, '*.xlsx'))):
        b = os.path.basename(f)
        if f not in taken and re.search(pattern, b, re.I): return f
    return None

SRC = sys.argv[1] if len(sys.argv) > 1 else pick(r'sale\s*2026')
if not SRC:
    sys.exit('ไม่พบไฟล์ยอดขายหลัก (ชื่อควรมี "Sale 2026") ในโฟลเดอร์ ' + ROOT)
PRODSRC = sys.argv[2] if len(sys.argv) > 2 else pick(r'product', taken=(SRC,))

OUT = os.path.join(ROOT, 'data', 'data.json')
print('อ่านยอดขายจาก:', os.path.basename(SRC))
print('อ่านสินค้าจาก:', os.path.basename(PRODSRC) if PRODSRC else '— (ไม่มี ข้ามแท็บสินค้า)')

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
MON = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']

def num(v):
    if v is None: return None
    if isinstance(v,(int,float)):
        if isinstance(v,float) and (math.isnan(v) or math.isinf(v)): return None
        return round(v,4)
    s=str(v).strip().replace(',','')
    if s in ('','-','#DIV/0!','#N/A',' '): return None
    try: return round(float(s),4)
    except: return None

out={}
out['meta']={'source':os.path.basename(SRC),'generated':datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}
if PRODSRC: out['meta']['productSource']=os.path.basename(PRODSRC)

# ---------- 1. YEARLY ----------
ws=wb['Sale Yealy']
rows=[list(r) for r in ws.iter_rows(max_row=89,max_col=7,values_only=True)]
yearly=collections.defaultdict(list)
cur=None
for r in rows:
    p=r[0]
    if p in (None,''): continue
    if p=='Platform': continue
    if p in ('Attribute','YoY diff'): continue
    y=num(r[1])
    if y is None: continue
    yearly[p].append({'year':int(y),'gmv':num(r[3]),'orders':num(r[4]),'aov':num(r[5]),'clicks':num(r[6])})
out['yearly']=dict(yearly)

# ---------- 2. MONTHLY 2026 + ADS ----------
# ชื่อชีตเปลี่ยนตามเวอร์ชันไฟล์ (v14 = '... + Ads Spend', v16 = '... + Ads 2026')
ADS_SHEET=None
for cand in ('Pivot Sale Monthly + Ads 2026','Pivot Sale Monthly + Ads Spend'):
    if cand in wb.sheetnames: ADS_SHEET=cand; break
if ADS_SHEET is None: sys.exit('ไม่พบชีตยอดขายรายเดือน + Ads')
ws=wb[ADS_SHEET]
g=[list(r) for r in ws.iter_rows(max_row=40,max_col=119,values_only=True)]
# หาคอลัมน์เริ่มของบล็อกแต่ละแพลตฟอร์มจากหัวตาราง (แถว index 3) แทนการ hardcode
hdr3=g[3]
hdr4=g[4]
def block_start(label):
    # ค้นเฉพาะฝั่งขวาของตาราง (คอลัมน์ 14 ขึ้นไป) เพราะฝั่งซ้ายเป็น Column Labels ของ pivot
    for ci in range(14,len(hdr3)):
        v=hdr3[ci]
        if v and str(v).strip()==label and str(hdr4[ci] or '').strip()=='GMV': return ci
    return None
BLK={'Lazada':block_start('Lazada'),'Shopee':block_start('Shopee'),'Tiktok':block_start('Tiktok')}
if None in BLK.values(): sys.exit('หาคอลัมน์บล็อกแพลตฟอร์มไม่เจอ: '+repr(BLK))
print('บล็อกแพลตฟอร์ม:', BLK)
plats=['Lazada','Line Shopping','Shopee','Tiktok','Facebook','Line Chat','Amaze','Website']
monthly=[]
for i,m in enumerate(MON):
    r=g[5+i]
    rec={'month':m,'byPlatform':{p:num(r[1+j]) for j,p in enumerate(plats)},
         'total':num(r[9]),'mom':num(r[10]),'adsSpend':num(r[11]),'adsPct':num(r[12]),'roi':num(r[13])}
    monthly.append(rec)
out['monthly2026']=monthly
# targets rows 19..29 (index)
tg={}
hdr=g[19]
for r in g[19:30]:
    if not r[0]: continue
    key=r[0]
    if key in ('Target',): continue
    tg[key]={MON[k]:num(r[1+k]) for k in range(12)}
    tg[key]['Total']=num(r[13])
out['targets2026']=tg

# platform ads detail (cols 15+)
def block(start, fields):
    res=[]
    for i,m in enumerate(MON):
        r=g[5+i]
        res.append({'month':m, **{f:num(r[start+k]) for k,f in enumerate(fields)}})
    return res
out['adsDetail']={
 'Lazada': block(BLK['Lazada'],['gmv','gmvTarget','adsMaxGmv','adsMaxSpend','roas','adsKwGmv','adsKwSpend','roasKw','adsAffGmv','adsAffSpend','roasAff','roi','diffMoMPct','diffMoM','diffTgPct','diffTg']),
 'Shopee': block(BLK['Shopee'],['gmv','gmvTarget','adsGmv','adsSpend','roas','affGmv','affSpend','roasAff','cpasGmv','cpasSpend','roasCpas','roi','diffMoMPct','diffMoM','diffTgPct','diffTg']),
 'Tiktok': block(BLK['Tiktok'],['gmv','gmvTarget','gmvMaxAllGmv','gmvMaxAllSpend','gmvMaxAllRoas','gmvMaxGmv','gmvMaxSpend','gmvMaxRoas','gmvMaxLiveGmv','gmvMaxLiveSpend','gmvMaxLiveRoas','liveRatio','cAdsSpend']),
}

# ---------- 3. DAILY from GMV Database ----------
ws=wb['GMV Database']
it=ws.iter_rows(values_only=True)
hdr=next(it)
daily=collections.defaultdict(lambda: collections.defaultdict(dict))
plat=year=mon=None
extra=collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(float)))
for r in it:
    if r[0]: plat=str(r[0]).strip()
    if r[1]: year=r[1]
    if r[2]: mon=str(r[2]).strip()
    d=num(r[4])
    if d is None or plat is None or year is None or mon is None: continue
    gmv=num(r[5])
    if gmv is None: continue
    key=f'{plat}|{year}|{mon}'
    daily[key][int(d)]={'gmv':gmv,'orders':num(r[6]),'aov':num(r[7]),'clicks':num(r[8]),'visitors':num(r[9]),'cvr':num(r[10]),
                        'cancelOrders':num(r[11]),'cancelGmv':num(r[12]),'refundOrders':num(r[13]),'refundGmv':num(r[14]),
                        'buyers':num(r[15]),'newBuyers':num(r[16]),'returnBuyers':num(r[17]),
                        'chProduct':num(r[22]),'chLive':num(r[23]),'chVideo':num(r[24]),'chPartner':num(r[25]),'chAds':num(r[26])}
out['daily']={k:{str(d):v for d,v in sorted(vv.items())} for k,vv in daily.items()}

# monthly aggregate 2025 vs 2026 for compare
agg=collections.defaultdict(lambda: collections.defaultdict(float))
for key,days in daily.items():
    p,y,m=key.split('|')
    for d,v in days.items():
        for f in ('gmv','orders','clicks','visitors','cancelGmv','newBuyers','returnBuyers','chProduct','chLive','chVideo','chPartner','chAds'):
            if v.get(f): agg[f'{p}|{y}|{m}'][f]+=v[f]
out['monthlyAgg']={k:{f:round(x,2) for f,x in v.items()} for k,v in agg.items()}

# ---------- 4. PRODUCTS (Mapping all stock) ----------
ws=wb['Mapping all stock']
rows=[list(r) for r in ws.iter_rows(min_row=3,max_row=74,max_col=50,values_only=True)]
hdr=rows[0]
# หาคอลัมน์จากหัวตารางแทน hardcode (ไฟล์เพิ่มคอลัมน์เดือนใหม่ทุกรอบ ตำแหน่งจึงเลื่อน)
MON_FULL={'SEP':'Sep','OCT':'Oct','NOV':'Nov','DEC':'Dec','JAN':'Jan','FEB':'Feb','MAR':'Mar','APR':'Apr','MAY':'May','JUN':'Jun','JUL':'Jul','AUG':'Aug'}
monthCols=[]
import re as _re
for ci,v in enumerate(hdr):
    if not v: continue
    s=' '.join(str(v).split()).strip()
    m=_re.match(r'^ยอดขาย\s*([A-Za-z]{3})\s*(\d{2})$', s)
    if m and m.group(1).upper() in MON_FULL:
        monthCols.append((ci, MON_FULL[m.group(1).upper()]+m.group(2)))
def find(label, frm=25):
    for ci in range(frm,len(hdr)):
        if hdr[ci] and ' '.join(str(hdr[ci]).split()).strip().startswith(label): return ci
    return None
C={'avg3m':find('Sale AVG 3 Month'),'avgDday':find('Sale AVG D-DAY'),
   'bauStatus':find('BAU สถานะ'),'ddayStatus':find('D-DAY'),
   'runScore':find('Run Score'),'stockLBL':None,'stockMHC':find('มหาชัย'),
   'planSend':find('จำนวน'),'status':find('Remark')}
# 'ลาดบัวหลวง' โผล่สองที่ (BAU สถานะ กับ ยอดสต็อก) — เอาตัวที่อยู่หลัง Run Score
rs=C['runScore']
C['stockLBL']=find('ลาดบัวหลวง', rs if rs else 25)
print('คอลัมน์สต็อก:', C, '| เดือน:', [m for _,m in monthCols])
prods=[]
for r in rows[2:]:
    if not r[0] or not r[2]: continue
    qty={name:num(r[ci]) for ci,name in monthCols}
    prods.append({'sku':str(r[0]).strip(),'category':str(r[1] or '').strip(),'name':str(r[2]).strip(),'pack':str(r[3] or '').strip(),
                  'qty':qty,'avg3m':num(r[C['avg3m']]),'avgDday':num(r[C['avgDday']]),
                  'bauStatus':str(r[C['bauStatus']] or '').strip(),'ddayStatus':str(r[C['ddayStatus']] or '').strip(),
                  'runScore':num(r[C['runScore']]),'stockLBL':num(r[C['stockLBL']]),'stockMHC':num(r[C['stockMHC']]),
                  'planSend':num(r[C['planSend']]) if C['planSend'] else None,
                  'status':str(r[C['status']] or '').strip() if C['status'] else ''})
out['products']=prods
out['mapMonths']=[m for _,m in monthCols]   # เดือนของชีต Mapping all stock (ของเดิม)

# price
ws=wb['Price']
prices=[]
for r in ws.iter_rows(min_row=5,max_row=104,max_col=18,values_only=True):
    if not r[1]: continue
    prices.append({'sku':str(r[1]).strip(),'type':str(r[2] or ''),'name':str(r[3] or ''),'pack':str(r[4] or ''),
                   'sticker':num(r[5]),'ao':num(r[6]),'mmpd':num(r[8]),'dd':num(r[10]),'mega':num(r[12]),'deepest':num(r[14])})
out['prices']=prices

# ---------- 5. AFFILIATE TIKTOK ----------
ws=wb['Affiliate รวม Database Tiktok']
it=ws.iter_rows(min_row=2,max_col=13,values_only=True)
aff=collections.defaultdict(lambda: collections.defaultdict(float))
creators=collections.defaultdict(lambda: collections.defaultdict(float))
namesByMonth=collections.defaultdict(set)
for r in it:
    m=r[0]; name=r[1]
    if not name: continue
    if isinstance(m,datetime.datetime): mk=m.strftime('%b %y')
    else: mk=str(m)[:12]
    gmv=num(r[2]) or 0; comm=num(r[11]) or 0; orders=num(r[4]) or 0
    aff[mk]['gmv']+=gmv; aff[mk]['commission']+=comm; aff[mk]['orders']+=orders; aff[mk]['creators']+=1
    aff[mk]['videos']+= (num(r[9]) or 0); aff[mk]['lives']+= (num(r[10]) or 0)
    creators[name]['gmv']+=gmv; creators[name]['commission']+=comm; creators[name]['orders']+=orders
    if gmv>0: namesByMonth[mk].add(name)
out['affTiktokMonthly']={k:{f:round(v,2) for f,v in d.items()} for k,d in aff.items()}
for k in out['affTiktokMonthly']: out['affTiktokMonthly'][k]['creators']=len(namesByMonth[k])
top=sorted(creators.items(), key=lambda kv:-kv[1]['gmv'])[:40]
out['affTiktokTop']=[{'name':k,**{f:round(v,2) for f,v in d.items()}} for k,d in top]

# followers
ws=wb['Affiliate List Tiktok']
out['affTiktokFollowers']={str(r[0]).strip():num(r[1]) for r in ws.iter_rows(min_row=2,max_col=2,values_only=True) if r[0]}

# shopee affiliate
ws=wb['Pivot Affiliate Shopee']
sa=[]
for r in ws.iter_rows(min_row=4,max_row=143,max_col=7,values_only=True):
    if not r[0] or str(r[0]).startswith('Grand'): continue
    sa.append({'name':str(r[0]).strip(),'gmv':num(r[1]),'qty':num(r[2]),'commission':num(r[3]),'roi':num(r[4]),'buyers':num(r[5]),'newBuyers':num(r[6])})
out['affShopee']=sa
ws=wb['Pivot Affiliate Shopee']
ch=[]
for r in ws.iter_rows(min_row=4,max_row=8,min_col=9,max_col=12,values_only=True):
    if r[0]: ch.append({'channel':str(r[0]),'gmv':num(r[1]),'qty':num(r[2]),'commission':num(r[3])})
out['affShopeeChannel']=ch

# lazada affiliate
ws=wb['Affiliate Database Lazada']
la=[]
for r in ws.iter_rows(min_row=1,max_row=32,max_col=7,values_only=True):
    la.append([str(c)[:40] if c is not None else '' for c in r])
out['affLazadaRaw']=la

# ---------- 6. REPURCHASE ----------
ws=wb['Shopee Report Repurchase']
rp=[]
for r in ws.iter_rows(min_row=5,max_row=7,max_col=16,values_only=True):
    if r[1]:
        rp.append({'metric':str(r[1]).strip(),**{MON[i]:num(r[2+i]) for i in range(12)},'total':num(r[14])})
out['shopeeRepurchase']=rp
ws=wb['Tiktok Report Repurchase']
tr=[]
for r in ws.iter_rows(min_row=6,max_row=37,max_col=14,values_only=True):
    if r[0] is None: continue
    tr.append({'day':num(r[0]),**{MON[i]:num(r[1+i]) for i in range(12)}})
out['tiktokRepurchaseDaily']=tr

# ---------- 7. D-DAY ----------
ws=wb['Sale D-Day']
dd=[]
for r in ws.iter_rows(min_row=2,max_row=41,max_col=13,values_only=True):
    if not r[2] or num(r[4]) is None: continue
    dd.append({'platform':str(r[0] or ''),'year':num(r[1]),'dday':str(r[2]),'gmv':num(r[4]),'orders':num(r[5]),'aov':num(r[6]),'clicks':num(r[7]),'visitors':num(r[8]),'cvr':num(r[9])})
out['dday']=dd

# ---------- 8. สินค้า & สต็อก (ไฟล์รายงานสินค้าแยกต่างหาก) ----------
def base_sku(s):
    """ตัดท้าย EXP…/GWP… ออก เพื่อยุบ lot ย่อยกลับเป็น SKU หลัก"""
    s=str(s).strip()
    s=re.sub(r'(EXP|GWP)[0-9A-Z]*$','',s)
    return s.strip() or str(s).strip()

if PRODSRC:
    pwb=openpyxl.load_workbook(PRODSRC, read_only=True, data_only=True)

    # 8.1 แคตตาล็อก: รูป + หมวด จากชีต pivot (คีย์เป็น SKU หลัก)
    catalog={}
    if '2026 Pivot Sale GMV + QTY' in pwb.sheetnames:
        for r in pwb['2026 Pivot Sale GMV + QTY'].iter_rows(min_row=6,max_col=3,values_only=True):
            if not r[0] or str(r[0]).startswith('Grand'): continue
            k=base_sku(r[0])
            img=str(r[1]) if r[1] and str(r[1]).startswith('http') else None
            cat=str(r[2]).strip() if r[2] and str(r[2])!='#N/A' else None
            e=catalog.setdefault(k,{'image':None,'category':None,'name':None})
            if img and not e['image']: e['image']=img
            if cat and not e['category']: e['category']=cat

    # 8.2 ยอดขายราย SKU รายเดือน (ก.ย. 68 – ปัจจุบัน) จากชีต Database
    MKEY=lambda d: MON[d.month-1]+str(d.year%100)
    sku=collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(float)))
    monthsSeen=set()
    for r in pwb['Database'].iter_rows(min_row=3,max_col=14,values_only=True):
        d=r[0]
        if not isinstance(d,datetime.datetime) or not r[4]: continue
        k=base_sku(r[4]); mk=MKEY(d); monthsSeen.add((d.year,d.month))
        cell=sku[k][mk]
        cell['gmv']+=num(r[10]) or 0        # ยอดขาย SKU Merchant
        cell['qty']+=num(r[11]) or 0        # การขาย SKU Merchant (ชิ้น)
        cell['orders']+=num(r[12]) or 0     # คำสั่งซื้อทั้งหมด
        cell['gmvNet']+=num(r[6]) or 0      # ยอดขายที่มีผล (หักยกเลิก/คืน)
        cell['qtyNet']+=num(r[8]) or 0
        e=catalog.setdefault(k,{'image':None,'category':None,'name':None})
        if not e['name'] and r[3] and str(r[3])!='#N/A': e['name']=str(r[3]).strip()
        if not e['category'] and r[2] and str(r[2])!='#N/A': e['category']=str(r[2]).strip()
    out['skuMonths']=[MON[m-1]+str(y%100) for y,m in sorted(monthsSeen)]
    out['skuMonthly']={k:{mk:{f:round(v,2) for f,v in d.items()} for mk,d in mm.items()} for k,mm in sku.items()}

    # 8.3 แยกแพลตฟอร์มราย SKU (ชีตนี้อัปเดตช้ากว่า — เก็บช่วงเดือนไว้บอกผู้ใช้ด้วย)
    plat=collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(float)))
    pm=set()
    if 'Database by Platform' in pwb.sheetnames:
        for r in pwb['Database by Platform'].iter_rows(min_row=2,max_col=15,values_only=True):
            d=r[0]
            if not isinstance(d,datetime.datetime) or not r[5] or not r[2]: continue
            pm.add((d.year,d.month))
            c=plat[base_sku(r[5])][str(r[2]).strip()]
            c['gmv']+=num(r[11]) or 0
            c['qty']+=num(r[12]) or 0
    out['skuPlatform']={k:{p:{f:round(v,2) for f,v in d.items()} for p,d in pp.items()} for k,pp in plat.items()}
    out['skuPlatformMonths']=[MON[m-1]+str(y%100) for y,m in sorted(pm)]

    # 8.4 สต็อกคงเหลือ (ชีต Stock) + วันที่อัปเดตจากหัวตาราง
    ws=pwb['Stock']
    sg=[list(r) for r in ws.iter_rows(max_row=ws.max_row,max_col=11,values_only=True)]
    asof=None
    for row in sg[:3]:
        for v in row:
            if v and 'Update' in str(v): asof=str(v).replace('Update','').strip(); break
        if asof: break
    out['stockAsOf']=asof
    stock=[]
    for r in sg[3:]:
        if not r[0] or not str(r[0]).strip(): continue
        stock.append({'sku':base_sku(r[0]),'category':str(r[1] or '').strip(),'brand':str(r[2] or '').strip(),
                      'type':str(r[3] or '').strip(),'lbl':num(r[5]) or 0,'mhc':num(r[6]) or 0,
                      'planQty':num(r[8]),'remark':str(r[9] or '').strip(),'status':str(r[10] or '').strip()})
    out['stock']=stock

    # 8.5 แผนของเข้า
    plan=[]
    if 'Plan in' in pwb.sheetnames:
        for r in pwb['Plan in'].iter_rows(min_row=2,max_col=7,values_only=True):
            if not r[1]: continue
            d=r[2] if isinstance(r[2],datetime.datetime) else None
            plan.append({'sku':base_sku(r[0]) if r[0] else None,'name':str(r[1]).strip(),
                         'date':d.strftime('%Y-%m-%d') if d else None,'qty':num(r[3]),
                         'unit':str(r[4] or '').strip(),'note':str(r[5] or '').strip()})
        plan.sort(key=lambda x: x['date'] or '')
    out['planIn']=plan

    # 8.6 ราคาขาย (ไฟล์สินค้ามีตารางราคาใหม่กว่าไฟล์หลัก)
    pr=[]
    for r in pwb['Price'].iter_rows(min_row=3,max_col=13,values_only=True):
        if not r[1]: continue
        pr.append({'sku':str(r[1]).strip(),'type':str(r[2] or '').strip(),'name':str(r[3] or '').strip(),
                   'pack':str(r[4] or '').strip(),'sticker':num(r[5]),'ao':num(r[6]),'mmpd':num(r[7]),
                   'dd':num(r[8]),'mega':num(r[9]),'deepest':num(r[10])})
    out['prices']=pr

    # 8.7 ชื่อ/สถานะเสริมจากไฟล์หลัก (ชีต Mapping all stock มี BAU / D-DAY / Run Score ที่ไฟล์สินค้าไม่มี)
    extra={}
    for p in out.get('products',[]):
        extra[base_sku(p['sku'])]={'bauStatus':p.get('bauStatus',''),'ddayStatus':p.get('ddayStatus',''),
                                   'runScore':p.get('runScore'),'avgDday':p.get('avgDday'),
                                   'name':p.get('name'),'pack':p.get('pack')}
    # ชื่อย่อสำหรับแสดงผล: เอาจากไฟล์หลักก่อน (สั้นและเป็นภาษาคน) แล้วค่อยตารางราคา
    # สุดท้ายค่อยใช้ชื่อยาวจาก Database ซึ่งเป็นชื่อเต็มบนหน้าร้าน
    byPrice={p['sku']:p for p in pr}
    for k,e in catalog.items():
        x=extra.get(k) or {}
        pz=byPrice.get(k)
        short=None
        if x.get('name'):
            short=x['name'] + ((' ' + x['pack']) if x.get('pack') and x['pack'] not in x['name'] else '')
        elif pz and pz.get('name'):
            short=pz['name'] + ((' ' + pz['pack']) if pz.get('pack') else '')
        elif e.get('name'):
            short=e['name'][:46]
        e['short']=(short or k).strip()
        if pz: e['pack']=pz.get('pack')
    out['skuExtra']=extra
    out['skuCatalog']=catalog

os.makedirs(os.path.dirname(OUT), exist_ok=True)
json.dump(out, open(OUT,'w'), ensure_ascii=False, separators=(',',':'))
print('เขียน:', OUT, os.path.getsize(OUT), 'bytes')
for k,v in out.items():
    print(k, type(v).__name__, len(v) if hasattr(v,'__len__') else '')
